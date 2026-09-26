#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dump_styles.py - print the values and the visible formatting of a workbook.

Handy for seeing exactly how a hand-formatted sample is styled, and for
comparing two files quickly.

    python3 dump_styles.py "BDB GSTR1 SYSTEM DATA JULY-2026.xlsx"
    python3 dump_styles.py DPBJUL.xlsx --values-only
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def describe(cell):
    bits = []
    font = cell.font
    if font.bold:
        bits.append("B")
    if font.underline:
        bits.append("U")
    if font.italic:
        bits.append("I")
    if font.name != "Calibri" or font.sz != 11:
        bits.append(f"{font.name}/{font.sz}")
    if cell.fill.fill_type:
        bits.append(f"fill:{cell.fill.fill_type}")
    if cell.alignment.horizontal:
        bits.append(cell.alignment.horizontal[0].upper())
    if cell.alignment.wrap_text:
        bits.append("wrap")
    if cell.alignment.vertical:
        bits.append(f"v:{cell.alignment.vertical[:1]}")
    if cell.number_format != "General":
        bits.append(f"[{cell.number_format}]")
    return "".join(bits) if bits else "-"


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    values_only = "--values-only" in sys.argv
    if not args:
        print(__doc__)
        return 2

    for name in args:
        path = Path(name)
        wb = load_workbook(path)
        print("=" * 78)
        print(f"{path.name}   sheets: {wb.sheetnames}")
        for ws in wb.worksheets:
            print(f"--- sheet '{ws.title}'  dims {ws.dimensions}  "
                  f"({ws.max_row} rows x {ws.max_column} cols)")
            for row in range(1, ws.max_row + 1):
                parts = []
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row, col)
                    if cell.value is None and values_only:
                        continue
                    value = cell.value
                    if isinstance(value, str):
                        value = value if len(value) <= 40 else value[:37] + "..."
                    if values_only:
                        parts.append(f"{cell.coordinate}={value!r}")
                    else:
                        parts.append(f"{cell.coordinate}={value!r}{describe(cell)}")
                if parts:
                    print(f"{row:>3} " + "  ".join(parts))
            widths = {}
            for dim in ws.column_dimensions.values():
                if dim.width:
                    for i in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
                        widths[get_column_letter(i)] = round(dim.width, 3)
            heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height}
            print(f"    widths : {widths}")
            print(f"    heights: {heights}")
            print(f"    default row height: {ws.sheet_format.defaultRowHeight}")
            print(f"    page: {ws.page_setup.orientation} paper={ws.page_setup.paperSize} "
                  f"scale={ws.page_setup.scale} margins="
                  f"({ws.page_margins.left:.3f},{ws.page_margins.right:.3f},"
                  f"{ws.page_margins.top:.3f},{ws.page_margins.bottom:.3f})")
            print(f"    freeze: {ws.freeze_panes}  merged: {ws.merged_cells.ranges}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
