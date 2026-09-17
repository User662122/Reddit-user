"""
Tests for E.py (AutoKeyPress).

Everything runs headless: no real keys are sent because the execution tests
use E.DRY_RUN, and the mouse watcher is fed fake mouse positions.

    pytest -q
"""

import datetime
import time

import openpyxl
import pytest

import E


# ── helpers ──────────────────────────────────────────────────────────────────
class FakeEvent:
    def __init__(self, name, event_type="down"):
        self.name = name
        self.event_type = event_type


class FakeKeyboard:
    """Queued stand-in for the `keyboard` module (no real key hooks)."""

    KEY_DOWN = "down"
    KEY_UP = "up"

    def __init__(self, events=(), default=None):
        self.events = list(events)
        self.default = default
        self.calls = 0

    def read_event(self, suppress=False):
        self.calls += 1
        if self.events:
            return self.events.pop(0)
        if self.default is not None:
            return self.default
        raise AssertionError("read_event() called with no queued key events")


@pytest.fixture(autouse=True)
def clean_state(monkeypatch):
    def reset():
        E.DRY_RUN_LOG.clear()
        E._WARNED.clear()
        E.reset_stop()
        E.DRY_RUN = False
        E.ACTION_GAP = 0.0
        E.TYPING_INTERVAL = 0.0
        E.TRIGGER_SETTLE = 0.0
        E.PASTE_COLUMNS = {"narration"}
        E.DATE_COLUMNS = {"date"}

    reset()
    monkeypatch.setattr(E, "keyboard", FakeKeyboard())
    monkeypatch.setattr(E, "HAS_KEYBOARD", True)
    yield
    reset()
    E.ACTION_GAP = 0.05
    E.TYPING_INTERVAL = 0.03
    E.TRIGGER_SETTLE = 0.25


def parse(script):
    return E.parse_script(script)


def make_workbook(path, extra_sheet=False):
    """3 data rows: mixed-case headers, one blank header, one blank row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["trigger", "keymap", "Amount", None, "Date"])
    ws.append(["cont", 'Enter + v"AMOUNT" + Tab + v"Date"', 1500.0, "x",
               datetime.datetime(2025, 1, 1)])
    ws.append(["Esc", 'v"NARRATION" + Enter', 250.25, "y",
               datetime.datetime(2025, 3, 5)])
    ws.append([None, None, None, None, None])          # blank -> skipped
    ws.append(["cont", 'v"AMOUNT"', 80.0, "z", "2025-12-31"])
    if extra_sheet:
        other = wb.create_sheet("Sheet2")
        other.append(["Trigger", "KeyMap"])
        other.append(["cont", "F5"])
    wb.save(path)
    return str(path)


# ── KeyMap parser ────────────────────────────────────────────────────────────
def test_parse_documented_example():
    assert parse('Enter + v"Date" + 2*Enter + v"NARRATION" + Tab + v"AMOUNT" + Enter') == [
        ("key", "enter"),
        ("type_col", "Date"),
        ("repeat", 2, ("key", "enter")),
        ("type_col", "NARRATION"),
        ("key", "tab"),
        ("type_col", "AMOUNT"),
        ("key", "enter"),
    ]


def test_parse_combo_single_modifier():
    assert parse("Ctrl+c") == [("combo", ["ctrl"], "c")]


def test_parse_combo_multiple_modifiers():
    assert parse("Ctrl+Shift+s") == [("combo", ["ctrl", "shift"], "s")]


def test_parse_combo_inside_a_longer_script():
    assert parse("Enter + Alt+F4 + Tab") == [
        ("key", "enter"),
        ("combo", ["alt"], "f4"),
        ("key", "tab"),
    ]


def test_parse_repeat_applies_to_a_combo():
    # "3*Ctrl+c" used to become (3 x Ctrl) followed by a bare "c"
    assert parse("3*Ctrl+c") == [("repeat", 3, ("combo", ["ctrl"], "c"))]


def test_parse_repeat_applies_to_columns_and_literals():
    assert parse('2*v"AMOUNT"') == [("repeat", 2, ("type_col", "AMOUNT"))]
    assert parse('2*t"ab"') == [("repeat", 2, ("type_literal", "ab"))]


def test_parse_literal_keeps_plus_signs():
    assert parse('t"a+b" + Enter') == [("type_literal", "a+b"), ("key", "enter")]


def test_parse_single_quotes_are_accepted():
    assert parse("v'AMOUNT'") == [("type_col", "AMOUNT")]


def test_parse_named_keys():
    assert parse("Esc") == [("key", "escape")]
    assert parse("PB") == [("key", "pause")]
    assert parse("IT") == [("key", "pause")]
    assert parse("F5") == [("key", "f5")]
    assert parse("Space") == [("key", "space")]
    assert parse("Up") == [("key", "up")]


def test_parse_single_characters():
    assert parse("a + 7 + .") == [("key", "a"), ("key", "7"), ("key", ".")]


def test_parse_skips_empty_tokens():
    assert parse("Enter++Tab+") == [("key", "enter"), ("key", "tab")]


def test_parse_zero_repeat_drops_the_action():
    assert parse("0*Enter") == []


def test_parse_lone_modifier_is_pressed_as_a_key():
    assert parse("Ctrl") == [("key", "ctrl")]


def test_parse_unknown_key_raises_instead_of_silently_dropping():
    with pytest.raises(E.ScriptError):
        parse("Frobnicate")


def test_parse_repeat_safety_limit():
    with pytest.raises(E.ScriptError):
        parse("99999*Enter")


def test_parse_modifier_needs_a_key():
    with pytest.raises(E.ScriptError):
        parse('Ctrl + v"AMOUNT"')


def test_split_actions_respects_quotes():
    assert E.split_actions('t"a+b" + Enter + Ctrl+c') == [
        't"a+b"', "Enter", "Ctrl", "c",
    ]


# ── Excel value formatting ───────────────────────────────────────────────────
def test_value_to_text_drops_the_trailing_dot_zero():
    # str(1500.0) == "1500.0" used to be typed straight into amount fields
    assert E.value_to_text(1500.0) == "1500"
    assert E.value_to_text(1500.5) == "1500.5"
    assert E.value_to_text(0.1 + 0.2) == "0.3"
    assert E.value_to_text(7) == "7"


def test_value_to_text_special_values():
    assert E.value_to_text(None) == ""
    assert E.value_to_text(True) == "TRUE"
    assert E.value_to_text(False) == "FALSE"
    assert E.value_to_text(datetime.date(2025, 1, 1)) == "01-01-2025"
    assert E.value_to_text(datetime.datetime(2025, 1, 1, 0, 0)) == "01-01-2025"
    assert E.value_to_text("  padded  ") == "padded"


@pytest.mark.parametrize("raw,expected", [
    (datetime.datetime(2025, 1, 1), "010125"),
    (datetime.date(2025, 3, 5), "050325"),
    ("01-01-2025", "010125"),
    ("5-3-2025", "050325"),
    ("01/01/25", "010125"),
    ("1.1.25", "010125"),
    ("2025-01-01", "010125"),           # ISO, year first
    ("2025-01-01 00:00:00", "010125"),  # stringified datetime
    ("01-Jan-2025", "010125"),
    ("Jan 5, 2025", "050125"),
    ("20250101", "010125"),             # text cell, YYYYMMDD
    ("010125", "010125"),               # already formatted
    (20250101, "010125"),               # numeric cell, YYYYMMDD
    (20250101.0, "010125"),
    ("2025-01-01T00:00:00", "010125"),
    (1500, "1500"),                     # not a date -> typed as-is
    ("13-13-2025", "13-13-2025"),       # impossible month -> unchanged
])
def test_format_date_ddmmyy(raw, expected):
    assert E.format_date_ddmmyy(raw) == expected


def test_format_date_never_raises_on_junk():
    for junk in ("", None, "not a date", "2025-13-45", "99/99/9999"):
        assert isinstance(E.format_date_ddmmyy(junk), str)


# ── Trigger resolution (regression: Esc never matched) ───────────────────────
def test_normalize_trigger():
    assert E.normalize_trigger("Esc") == ("key", "escape")
    assert E.normalize_trigger("cont") == ("cont", None)
    assert E.normalize_trigger("CONT") == ("cont", None)
    assert E.normalize_trigger("2211") == ("digits", "2211")
    assert E.normalize_trigger("PB") == ("key", "pause")
    assert E.normalize_trigger("IT") == ("key", "pause")
    assert E.normalize_trigger("F5") == ("key", "f5")
    assert E.normalize_trigger("") == ("none", None)


def test_esc_trigger_accepts_the_keyboard_module_spelling():
    # the keyboard module reports Escape as "esc" while KEY_MAP maps Esc ->
    # "escape"; comparing those two literally meant waiting forever.
    assert "esc" in E.accepted_trigger_names("escape")
    assert "escape" in E.accepted_trigger_names("escape")


def test_pause_trigger_accepts_every_spelling():
    names = E.accepted_trigger_names("pause")
    for spelling in ("pause", "break", "pb", "it"):
        assert spelling in names


def test_digit_buffer_matches_a_sequence():
    buffer, matched = "", False
    for key in ("2", "2", "1"):
        buffer, matched = E.digit_buffer_step(buffer, key, "2211")
        assert matched is False
    buffer, matched = E.digit_buffer_step(buffer, "1", "2211")
    assert matched is True
    assert buffer == ""


def test_digit_buffer_resets_on_other_keys():
    buffer, _ = E.digit_buffer_step("", "2", "2211")
    assert buffer == "2"
    assert E.digit_buffer_step(buffer, "a", "2211") == ("", False)


def test_digit_buffer_accepts_numpad_names():
    buffer, matched = "", False
    for key in ("num 2", "num 2", "num 1", "num 1"):
        buffer, matched = E.digit_buffer_step(buffer, key, "2211")
    assert matched is True


# ── Excel loading ────────────────────────────────────────────────────────────
def test_load_excel_reads_headers_and_rows(tmp_path):
    sheet, label = E.load_excel(make_workbook(tmp_path / "b.xlsx"))
    assert label == "Sheet1"
    assert sheet.headers[:3] == ["trigger", "keymap", "Amount"]
    assert sheet.headers[3].startswith("Column")   # blank header got a name
    assert len(sheet.rows) == 3                    # blank row skipped
    assert [r["__excel_row__"] for r in sheet.rows] == [2, 3, 5]


def test_load_excel_is_case_insensitive(tmp_path):
    sheet, _ = E.load_excel(make_workbook(tmp_path / "b.xlsx"))
    assert sheet.has("Trigger") and sheet.has("TRIGGER") and sheet.has("trigger")
    assert sheet.get(sheet.rows[0], "Trigger") == "cont"
    assert sheet.get(sheet.rows[0], "amount") == 1500.0
    assert not sheet.has("Nope")
    assert sheet.get(sheet.rows[0], "Nope", "fallback") == "fallback"


def test_load_excel_sheet_argument_is_honoured(tmp_path):
    path = make_workbook(tmp_path / "b.xlsx", extra_sheet=True)
    sheet, label = E.load_excel(path, "Sheet2")
    assert label == "Sheet2"
    assert len(sheet.rows) == 1
    with pytest.raises(E.ScriptError):
        E.load_excel(path, "DoesNotExist")


def test_load_excel_does_not_type_formula_text(tmp_path):
    """Formula cells must never be typed as '=C2*2' (data_only=True)."""
    path = tmp_path / "f.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Trigger", "KeyMap", "Amount"])
    ws.append(["cont", 'v"AMOUNT"', 40])
    ws.append(["cont", 'v"AMOUNT"', "=C2*2"])
    wb.save(path)

    sheet, _ = E.load_excel(str(path))
    formula_value = sheet.get(sheet.rows[1], "Amount")
    # openpyxl cannot compute the formula, so the cached value is missing ->
    # empty. What matters is that the literal text is never sent.
    assert formula_value != "=C2*2"
    assert E.value_to_text(formula_value) == ""


def test_load_excel_missing_file(tmp_path):
    with pytest.raises(E.ScriptError):
        E.load_excel(str(tmp_path / "nope.xlsx"))


# ── Execution (dry run) ──────────────────────────────────────────────────────
def test_execute_dry_run_types_the_expected_things():
    E.DRY_RUN = True
    sheet = E.SheetData(["Trigger", "KeyMap", "AMOUNT", "Date", "NARRATION"], [{}])
    row = {
        "AMOUNT": 1500.0,
        "Date": datetime.datetime(2025, 1, 1),
        "NARRATION": "Cash paid",
    }
    actions = parse('Enter + v"Date" + 2*Enter + v"NARRATION" + Tab + v"AMOUNT"')
    E.execute_actions(actions, row, sheet)

    assert ("type", "010125") in E.DRY_RUN_LOG       # Date column -> DDMMYY
    assert ("type", "1500") in E.DRY_RUN_LOG         # 1500.0 -> "1500"
    assert ("type", "Cash paid") in E.DRY_RUN_LOG or \
           ("paste", "Cash paid") in E.DRY_RUN_LOG
    assert E.DRY_RUN_LOG.count(("key", "enter")) == 3   # 1 + 2*Enter
    assert ("key", "tab") in E.DRY_RUN_LOG


def test_execute_repeat_combo_presses_the_whole_combo():
    E.DRY_RUN = True
    E.execute_actions(parse("2*Ctrl+c"), {}, None)
    assert E.DRY_RUN_LOG == [
        ("combo", ["ctrl"], "c"),
        ("combo", ["ctrl"], "c"),
    ]


def test_execute_skips_missing_columns_with_a_warning(capsys):
    E.DRY_RUN = True
    E.execute_actions(parse('v"NOT_THERE" + Enter'), {}, None)
    assert ("key", "enter") in E.DRY_RUN_LOG
    assert "NOT_THERE" in capsys.readouterr().out


def test_unicode_text_goes_through_the_clipboard():
    E.DRY_RUN = True
    assert E._needs_clipboard("Café \u20b9")
    assert not E._needs_clipboard("plain 123")
    E.execute_actions(parse('t"Rupees \u20b9100"'), {}, None)
    assert ("paste", "Rupees \u20b9100") in E.DRY_RUN_LOG or \
           ("type", "Rupees \u20b9100") in E.DRY_RUN_LOG


# ── Mouse stop ───────────────────────────────────────────────────────────────
def test_check_stop_raises_only_when_armed():
    E.DRY_RUN = True
    E.execute_actions(parse("Enter"), {}, None)          # no stop -> runs
    E.request_stop("mouse moved 40px")
    with pytest.raises(E.AutomationStopped):
        E.check_stop()
    with pytest.raises(E.AutomationStopped):
        E.execute_actions(parse("Enter"), {}, None)


def test_mouse_watcher_stops_on_movement(monkeypatch):
    positions = iter([(100, 100), (100, 100), (400, 400), (400, 400)])
    monkeypatch.setattr(E, "mouse_position", lambda: next(positions, (400, 400)))
    watcher = E.MouseStopWatcher(threshold=5, poll=0.005)
    watcher.watch()
    deadline = time.time() + 2
    while not E.STOP_EVENT.is_set() and time.time() < deadline:
        time.sleep(0.01)
    watcher.cancel()
    assert E.STOP_EVENT.is_set()
    assert "mouse moved" in E.STOP_REASON


def test_mouse_watcher_ignores_jitter(monkeypatch):
    monkeypatch.setattr(E, "mouse_position", lambda: (100, 101))
    watcher = E.MouseStopWatcher(threshold=5, poll=0.005)
    watcher.watch()
    time.sleep(0.1)
    watcher.cancel()
    assert not E.STOP_EVENT.is_set()


def test_mouse_stop_can_be_disabled():
    watcher = E.MouseStopWatcher(threshold=0)
    watcher.watch()          # threshold <= 0 -> never arms
    assert not watcher.is_alive()
    assert not E.STOP_EVENT.is_set()


def test_running_automation_aborts_when_the_mouse_moves(monkeypatch):
    """End-to-end: real watcher thread + real executor = stop mid-row."""
    E.DRY_RUN = True
    E.ACTION_GAP = 0.01
    calls = {"n": 0}

    def fake_position():
        calls["n"] += 1
        return (100, 100) if calls["n"] == 1 else (900, 900)

    monkeypatch.setattr(E, "mouse_position", fake_position)

    actions = parse(" + ".join(["Enter"] * 200))   # ~2s of work
    watcher = E.MouseStopWatcher(threshold=5, poll=0.005)
    watcher.watch()
    try:
        with pytest.raises(E.AutomationStopped):
            E.execute_actions(actions, {}, None)
    finally:
        watcher.cancel()

    pressed = E.DRY_RUN_LOG.count(("key", "enter"))
    assert pressed < 200          # the row was cut short, not completed
    assert "mouse moved" in E.STOP_REASON


# ── main() ───────────────────────────────────────────────────────────────────
def test_main_dry_run_end_to_end(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("builtins.input", lambda *a, **k: "")
    E.keyboard.default = FakeEvent("esc")        # for the Esc-triggered row
    code = E.main([make_workbook(tmp_path / "b.xlsx"), "--dry-run", "--delay", "0"])
    out = capsys.readouterr().out
    assert code == 0
    assert "All done - 3/3 row(s) processed." in out
    assert ("type", "1500") in E.DRY_RUN_LOG
    assert ("type", "010125") in E.DRY_RUN_LOG


def test_main_uses_the_selected_sheet(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("builtins.input", lambda *a, **k: "")
    path = make_workbook(tmp_path / "b.xlsx", extra_sheet=True)
    code = E.main([path, "--sheet", "Sheet2", "--dry-run", "--delay", "0"])
    out = capsys.readouterr().out
    assert code == 0
    assert "Sheet : Sheet2" in out
    assert ("key", "f5") in E.DRY_RUN_LOG


def test_main_rejects_a_bad_keymap_before_typing_anything(tmp_path, capsys):
    path = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Trigger", "KeyMap"])
    ws.append(["cont", "Enter + Frobnicate"])
    wb.save(path)

    code = E.main([str(path), "--dry-run"])
    out = capsys.readouterr().out
    assert code == 1
    assert "Unknown key 'Frobnicate'" in out
    assert E.DRY_RUN_LOG == []          # nothing was typed


def test_main_reports_a_missing_file(tmp_path, capsys):
    code = E.main([str(tmp_path / "nope.xlsx"), "--dry-run"])
    assert code == 1
    assert "ERROR" in capsys.readouterr().out


def test_main_reports_missing_required_columns(tmp_path, capsys):
    path = tmp_path / "cols.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Foo", "Bar"])
    wb.save(path)
    code = E.main([str(path), "--dry-run"])
    assert code == 1
    assert "Trigger" in capsys.readouterr().out


def test_main_stop_during_a_row(tmp_path, monkeypatch, capsys):
    """A mouse stop mid-run ends the run and reports how much was done."""
    path = make_workbook(tmp_path / "b.xlsx")
    real_execute = E.execute_actions

    def stopping_execute(actions, row, sheet=None):
        real_execute(actions[:1], row, sheet)
        E.request_stop("mouse moved 12px (threshold 5px)")
        E.check_stop()

    monkeypatch.setattr(E, "execute_actions", stopping_execute)
    monkeypatch.setattr("builtins.input", lambda *a, **k: "n")
    code = E.main([path, "--dry-run", "--delay", "0"])
    out = capsys.readouterr().out
    assert code == 2
    assert "STOPPED" in out
    assert "Completed 0/3 row(s)" in out
    assert "No keys are held down." in out


def test_countdown_honours_fractional_seconds():
    start = time.time()
    E._countdown(0.05)
    assert time.time() - start >= 0.04


# ── real (non-dry-run) key sending, with a fake pyautogui ────────────────────
class FakePyAutoGUI:
    """Records what the real sending code would have done."""

    def __init__(self):
        self.events = []
        self.held = []

    def press(self, key):
        self.events.append(("press", key))

    def keyDown(self, key):
        self.held.append(key)
        self.events.append(("down", key))

    def keyUp(self, key):
        if key in self.held:
            self.held.remove(key)
        self.events.append(("up", key))

    def position(self):
        return (0, 0)


@pytest.fixture
def fake_gui(monkeypatch):
    fake = FakePyAutoGUI()
    monkeypatch.setattr(E, "pyautogui", fake)
    monkeypatch.setattr(E, "HAS_PYAUTOGUI", True)
    return fake


def presses(fake):
    return [e for e in fake.events if e[0] == "press"]


def test_combo_releases_its_modifiers(fake_gui):
    E.execute_actions(parse("Ctrl+Shift+s + Enter"), {}, None)
    assert ("down", "ctrl") in fake_gui.events
    assert ("down", "shift") in fake_gui.events
    assert ("press", "s") in fake_gui.events
    assert fake_gui.held == []          # nothing left held down
    assert ("press", "enter") in fake_gui.events


def test_typing_stops_immediately_when_the_mouse_moves(fake_gui, monkeypatch):
    def press(key):
        fake_gui.events.append(("press", key))
        if len(presses(fake_gui)) >= 5:
            E.request_stop("mouse moved 30px (threshold 5px)")

    monkeypatch.setattr(fake_gui, "press", press)
    with pytest.raises(E.AutomationStopped):
        E.execute_actions(parse('t"0123456789" + Ctrl+c'), {}, None)
    assert len(presses(fake_gui)) == 5      # the rest of the row was not typed
    assert fake_gui.held == []


def test_stop_in_the_middle_of_a_combo_releases_the_modifier(fake_gui, monkeypatch):
    def press(key):
        fake_gui.events.append(("press", key))
        if key == "c":
            raise E.AutomationStopped("mouse moved 30px")

    monkeypatch.setattr(fake_gui, "press", press)
    with pytest.raises(E.AutomationStopped):
        E.execute_actions(parse("Ctrl+c + Enter"), {}, None)
    assert fake_gui.held == []              # Ctrl was released, not stuck
    assert ("press", "enter") not in fake_gui.events


def test_paste_falls_back_to_typing_without_pyperclip(fake_gui, monkeypatch):
    monkeypatch.setattr(E, "HAS_PYPERCLIP", False)
    monkeypatch.setattr(E, "pyperclip", None)
    E.execute_actions(parse('v"NARRATION"'), {"NARRATION": "Cash paid"}, None)
    assert "".join(k for kind, k in fake_gui.events if kind == "press") == "Cash paid"


def test_numeric_trigger_cell_stays_a_digit_sequence(tmp_path):
    """2211 stored as a number must not become the text '2211.0'."""
    path = tmp_path / "num.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Trigger", "KeyMap", "Amount"])
    ws.append([2211.0, 'v"AMOUNT"', 5])
    wb.save(path)

    sheet, _ = E.load_excel(str(path))
    trigger = E.value_to_text(sheet.get(sheet.rows[0], "Trigger")).strip()
    assert trigger == "2211"
    assert E.normalize_trigger(trigger) == ("digits", "2211")


@pytest.mark.skipif(E.HAS_PYAUTOGUI, reason="needs a machine without pyautogui")
def test_main_fails_fast_when_pyautogui_is_missing(tmp_path, capsys):
    code = E.main([make_workbook(tmp_path / "b.xlsx")])
    out = capsys.readouterr().out
    assert code == 1
    assert "pyautogui could not be imported" in out
    assert E.DRY_RUN_LOG == []


def test_run_rows_rearms_the_watcher_for_every_row(monkeypatch, capsys):
    """Regression: one watcher reused across rows crashed on row 2."""
    monkeypatch.setattr("builtins.input", lambda *a, **k: "")
    monkeypatch.setattr(E, "mouse_position", lambda: (100, 100))
    args = E.build_arg_parser().parse_args(["--delay", "0"])
    E.DRY_RUN = True

    rows = [{"Trigger": "cont", "KeyMap": "Enter + Enter"} for _ in range(3)]
    sheet = E.SheetData(["Trigger", "KeyMap"], rows)
    prepared = [(("cont", None), E.parse_script("Enter + Enter"), "cont",
                 "Enter + Enter") for _ in rows]

    code, completed, stopped = E.run_rows(args, sheet, prepared, 0, 0,
                                          mouse_stop=True)
    assert (code, completed, stopped) == (0, 3, None)
    assert E.DRY_RUN_LOG.count(("key", "enter")) == 6
    assert "threads can only be started once" not in capsys.readouterr().out


def test_watcher_thread_is_stopped_by_shutdown(monkeypatch):
    monkeypatch.setattr(E, "mouse_position", lambda: (100, 100))
    watcher = E.MouseStopWatcher(threshold=5, poll=0.005)
    watcher.watch()
    assert watcher.is_alive()
    watcher.cancel()
    watcher.shutdown()
    deadline = time.time() + 1
    while watcher.is_alive() and time.time() < deadline:
        time.sleep(0.01)
    assert not watcher.is_alive()


# ── trigger waiting (real code path, fake keyboard events) ───────────────────
def test_wait_for_trigger_key_accepts_the_keyboard_spelling_of_esc():
    """Regression: 'Esc' waited forever because esc != escape."""
    E.keyboard.events = [FakeEvent("shift", "down"), FakeEvent("esc", "down")]
    E.wait_for_trigger_key("escape")            # returns instead of hanging
    assert E.keyboard.calls == 2


def test_wait_for_trigger_key_ignores_key_up_and_other_keys():
    E.keyboard.events = [
        FakeEvent("a", "down"),
        FakeEvent("a", "up"),
        FakeEvent("ctrl", "down"),
        FakeEvent("f5", "down"),
    ]
    E.wait_for_trigger_key("f5")
    assert E.keyboard.calls == 4


def test_wait_for_trigger_key_accepts_pb_and_it_as_pause():
    for trigger in ("PB", "IT"):
        key = E.normalize_trigger(trigger)[1]
        E.keyboard.events = [FakeEvent("break", "down")]
        E.wait_for_trigger_key(key)


def test_wait_for_trigger_digits_needs_the_whole_sequence():
    E.keyboard.events = [
        FakeEvent("2", "down"), FakeEvent("2", "down"), FakeEvent("1", "down"),
        FakeEvent("a", "down"),                                   # resets
        FakeEvent("2", "down"), FakeEvent("2", "down"),
        FakeEvent("1", "down"), FakeEvent("1", "down"),
    ]
    E.wait_for_trigger_digits("2211")
    assert E.keyboard.calls == 8


def test_wait_for_trigger_uses_the_enter_fallback_without_keyboard(monkeypatch):
    monkeypatch.setattr(E, "HAS_KEYBOARD", False)
    seen = []
    monkeypatch.setattr("builtins.input", lambda prompt="": seen.append(prompt))
    E.wait_for_trigger(("key", "escape"))
    assert seen and "ESCAPE" in seen[0]


def test_keyboard_errors_are_reported_not_swallowed(monkeypatch):
    class Broken:
        KEY_DOWN = "down"

        def read_event(self, suppress=False):
            raise OSError("could not listen to the keyboard")

    monkeypatch.setattr(E, "keyboard", Broken())
    monkeypatch.setattr(E.time, "sleep", lambda s: None)
    with pytest.raises(E.ScriptError) as exc:
        E.wait_for_trigger_key("escape")
    assert "keyboard module keeps failing" in str(exc.value)
